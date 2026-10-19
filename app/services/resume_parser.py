import io
from pathlib import Path
from typing import Tuple, Optional
import fitz  # PyMuPDF
from docx import Document

from app.utils.text_cleaner import clean_text
from app.utils.file_validator import validate_resume_file
from app.utils.hash_utils import compute_file_hash, compute_text_hash

class ResumeParserService:
    def __init__(self, max_size_mb: int = 15):
        self.max_size_mb = max_size_mb

    def extract_text_from_pdf(self, file_path: Path | str) -> str:
        """Extracts text from PDF using PyMuPDF (fitz) with layout-aware blocks sorted in reading order."""
        doc = fitz.open(str(file_path))
        extracted_pages = []
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            blocks = page.get_text("blocks") or []
            # Sort blocks primarily by y0 (vertical coordinate snapped by line height), secondarily by x0
            sorted_blocks = sorted(blocks, key=lambda b: (round(b[1] / 12.0) * 12.0, b[0]))
            page_text = "\n".join(b[4].strip() for b in sorted_blocks if len(b) > 4 and b[4].strip())
            extracted_pages.append(page_text)
        doc.close()
        return "\n".join(extracted_pages)

    def extract_text_from_docx(self, file_path: Path | str) -> str:
        """Extracts text from DOCX files including paragraphs and table contents."""
        doc = Document(str(file_path))
        text_parts = []
        
        # Paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
                
        # Tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_parts.append(" | ".join(row_texts))
                    
        return "\n".join(text_parts)

    def parse_file(self, file_path: Path | str) -> Tuple[bool, Optional[str], Optional[str], Optional[str], Optional[str]]:
        """
        Validates and parses a resume file (PDF, DOCX, TXT).
        Returns: (success, raw_text, cleaned_text, file_hash, error_message)
        """
        path = Path(file_path)
        is_valid, error = validate_resume_file(path, max_size_mb=self.max_size_mb)
        if not is_valid:
            return False, None, None, None, error

        try:
            file_hash = compute_file_hash(path)
            suffix = path.suffix.lower()

            if suffix == ".pdf":
                raw_text = self.extract_text_from_pdf(path)
            elif suffix == ".docx":
                raw_text = self.extract_text_from_docx(path)
            elif suffix == ".txt":
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    raw_text = f.read()
            else:
                return False, None, None, None, f"Unsupported format: {suffix}"

            if not raw_text.strip():
                return False, None, None, file_hash, "Extracted text is empty (scanned image or empty document)"

            cleaned_text = clean_text(raw_text)
            return True, raw_text, cleaned_text, file_hash, None

        except Exception as e:
            return False, None, None, None, f"Parsing exception: {str(e)}"
